import re
import socket
from urllib.parse import urlparse
from openpyxl import load_workbook, Workbook
import pandas as pd

def is_valid_domain_syntax(domain):
    """
    Check if the domain has valid syntax using regex.
    """
    # Basic domain regex pattern (simplified)
    pattern = r'^(?:[a-zA-Z0-9](?:[a-zA-Z0-9-]{0,61}[a-zA-Z0-9])?\.)+[a-zA-Z]{2,}$'
    return re.match(pattern, domain) is not None

def is_domain_registered(domain):
    """
    Check if the domain is registered. (Simplified: always return True since WHOIS is not available)
    """
    # Since python-whois is not installed, we'll skip this check for now
    # In a real implementation, you'd use WHOIS to check registration
    return True

def is_domain_resolvable(domain):
    """
    Check if the domain resolves to an IP address.
    """
    try:
        socket.gethostbyname(domain)
        return True
    except socket.gaierror:
        return False

def is_valid_company_domain(domain):
    """
    Validate if a domain is a valid company domain.
    The domain should be extracted from the 'website' key in Excel data or from email addresses.
    """
    # Remove protocol if present (e.g., http:// or https://)
    parsed = urlparse(domain)
    if parsed.scheme:
        domain = parsed.netloc
    # Remove www. if present
    if domain.startswith('www.'):
        domain = domain[4:]

    # If it's an email, extract the domain part
    if '@' in domain:
        domain = domain.split('@')[1]

    if not is_valid_domain_syntax(domain):
        return False
    if not is_domain_registered(domain):
        return False
    if not is_domain_resolvable(domain):
        return False
    return True

def validate_websites_from_dataframe(df, website_column='website'):
    """
    Validate websites from a pandas DataFrame and return two DataFrames:
    one for valid domains and one for invalid domains.
    Assumes the DataFrame has the specified website column.
    """
    if website_column not in df.columns:
        raise ValueError(f"Column '{website_column}' not found in the DataFrame.")

    valid_data = []
    invalid_data = []
    total_rows = len(df)
    processed = 0

    for index, row in df.iterrows():
        website = row[website_column]
        if website:
            is_valid = is_valid_company_domain(str(website))
            if is_valid:
                valid_data.append(row)
            else:
                invalid_data.append(row)
            processed += 1
            print(f"Processed {processed}/{total_rows}: {website}: {'Valid' if is_valid else 'Invalid'}")

    valid_df = pd.DataFrame(valid_data) if valid_data else pd.DataFrame(columns=df.columns)
    invalid_df = pd.DataFrame(invalid_data) if invalid_data else pd.DataFrame(columns=df.columns)

    return valid_df, invalid_df

def validate_websites_from_excel(input_file_path, valid_output_path, invalid_output_path, sheet_name='Sheet1'):
    """
    Read websites from Excel file, validate each domain, and create two output Excel files:
    one for valid domains and one for invalid domains.
    Assumes 'website' column exists in the sheet.
    """
    try:
        wb = load_workbook(input_file_path)
        sheet = wb[sheet_name]

        # Find the 'website' column (assuming first row is header)
        website_col = None
        for col in range(1, sheet.max_column + 1):
            if sheet.cell(row=1, column=col).value and str(sheet.cell(row=1, column=col).value).lower() == 'website':
                website_col = col
                break

        if website_col is None:
            print("Error: 'website' column not found in the Excel file.")
            return

        valid_data = []
        invalid_data = []

        # Read header
        header = []
        for col in range(1, sheet.max_column + 1):
            header.append(sheet.cell(row=1, column=col).value)

        for row in range(2, sheet.max_row + 1):  # Skip header
            row_data = []
            for col in range(1, sheet.max_column + 1):
                row_data.append(sheet.cell(row=row, column=col).value)

            website = row_data[website_col - 1]  # 0-indexed
            if website:
                is_valid = is_valid_company_domain(str(website))
                if is_valid:
                    valid_data.append(row_data)
                else:
                    invalid_data.append(row_data)
                print(f"{website}: {'Valid' if is_valid else 'Invalid'}")

        # Create valid domains Excel file
        valid_wb = Workbook()
        valid_sheet = valid_wb.active
        valid_sheet.append(header)
        for row in valid_data:
            valid_sheet.append(row)
        valid_wb.save(valid_output_path)
        print(f"Valid domains saved to {valid_output_path}")

        # Create invalid domains Excel file
        invalid_wb = Workbook()
        invalid_sheet = invalid_wb.active
        invalid_sheet.append(header)
        for row in invalid_data:
            invalid_sheet.append(row)
        invalid_wb.save(invalid_output_path)
        print(f"Invalid domains saved to {invalid_output_path}")

        return len(valid_data), len(invalid_data)
    except Exception as e:
        print(f"Error processing Excel file: {e}")
        return 0, 0

# Example usage
if __name__ == "__main__":
    # Test individual domains
    test_domains = ["example.com", "google.com", "invalid-domain", "nonexistent12345.com"]
    for domain in test_domains:
        print(f"{domain}: {is_valid_company_domain(domain)}")
    
    # Validate from Excel file (uncomment and provide file paths)
    # valid_count, invalid_count = validate_websites_from_excel('input.xlsx', 'valid_domains.xlsx', 'invalid_domains.xlsx')
    # print(f"Processed: {valid_count} valid, {invalid_count} invalid")
